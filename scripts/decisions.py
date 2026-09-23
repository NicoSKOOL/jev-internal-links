#!/usr/bin/env python3
"""Step 7: one row per Jev decision, with what happened to it and why.

Joins judgements (what Jev said), config thresholds (what the rules did),
allocation + cap log (what the caps did) and links (what got an anchor).
Optionally compares against an earlier run to show run-to-run consistency.

Usage: python decisions.py [--compare runs/<older-run>/data/judgements.json]
Output: out/jev-decisions.csv   every decision, flat, spreadsheet-ready
        out/jev-decisions.xlsx  same, styled, plus a Linked sheet to approve
        out/jev-dashboard.html  interactive dashboard (no server needed)
"""
import argparse
import csv
import json
import statistics
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

import yaml

from _env import WORKDIR

HERE = WORKDIR  # data/, cache/, out/ and config.yaml live where you run it
DATA = HERE / "data"
OUT = HERE / "out"
TEMPLATE = Path(__file__).parent / "dashboard_template.html"
CFG = yaml.safe_load((HERE / "config.yaml").read_text())

# Outcome codes, in pipeline order. Labels are what a reader sees.
OUTCOMES = {
    "linked_apply": "Linked, high confidence",
    "linked_review": "Linked, needs your review",
    "no_anchor": "Kept, but no clean anchor in the paragraph",
    "source_full": "Source page already got its new links",
    "too_close": "Too close to another new link",
    "target_full": "Target page already got enough new links",
    "duplicate_pair": "Page already links to that target elsewhere",
    "model_said_none": "Jev said none of these pages fit",
    "low_confidence": "Jev was not sure enough",
    "not_warranted": "Paragraph doesn't need a link",
    "no_natural_anchor": "No natural anchor phrase",
    "self_link": "Would link to itself",
    "api_error": "API error",
}


def path(url):
    return urlparse(url).path or "/"


def classify(ans, cand, th):
    """Replay allocate.py's filters to name the first rule that stopped it."""
    if not ans:
        return "api_error"
    bt = ans["best_target"]
    if bt["choice"] == "none":
        return "model_said_none"
    if bt["confidence"] < th["min_confidence"]:
        return "low_confidence"
    if ans["link_warranted"]["noul"] < th["min_link_warranted"]:
        return "not_warranted"
    if ans["anchor_available"]["noul"] < th["min_anchor_available"]:
        return "no_natural_anchor"
    if cand["targets"][int(bt["choice"]) - 1]["canon"] == cand["source_canon"]:
        return "self_link"
    return None


def build_rows(compare_path=None):
    judgements = json.loads((DATA / "judgements.json").read_text())
    alloc = {(c["source_canon"], c["p_i"]): c
             for c in json.loads((DATA / "allocation.json").read_text())}
    links = {(c["source_canon"], c["p_i"]): c
             for c in json.loads((DATA / "links.json").read_text())}
    caplog = json.loads((DATA / "cap_log.json").read_text())
    th = CFG["thresholds"]

    prev = {}
    if compare_path:
        for r in json.loads(Path(compare_path).read_text()):
            a = r.get("response", {}).get("answers")
            if a:
                prev[(r["candidate"]["source_canon"], r["candidate"]["p_i"])] = a

    rows = []
    for n, r in enumerate(judgements, 1):
        c, resp = r["candidate"], r.get("response", {})
        ans = resp.get("answers")
        key = (c["source_canon"], c["p_i"])
        outcome = classify(ans, c, th)

        chosen = None
        if ans and ans["best_target"]["choice"] != "none":
            chosen = c["targets"][int(ans["best_target"]["choice"]) - 1]

        if outcome is None:
            if key in links:
                outcome = "linked_" + links[key]["action"]
            elif key in alloc:
                outcome = "no_anchor"
            else:
                outcome = caplog.get(
                    f"{key[0]}|{key[1]}|{chosen['canon']}", "source_full")

        probs = ans["best_target"]["probabilities"] if ans else {}
        cands = [{
            "url": t["url"], "title": t["title"],
            "prob": probs.get(str(i), 0), "sim": t["sim"],
            "inbound": t["inbound_links"],
        } for i, t in enumerate(c["targets"], 1)]

        link = links.get(key) or alloc.get(key) or {}
        p = prev.get(key)
        row = {
            "id": n,
            "outcome": outcome,
            "outcome_label": OUTCOMES[outcome],
            "source_url": c["source_url"],
            "source_title": c["source_title"],
            "paragraph_no": c["p_i"],
            "paragraph": c["text"],
            "jev_choice": ans["best_target"]["choice"] if ans else "",
            "chosen_url": chosen["url"] if chosen else "",
            "chosen_title": chosen["title"] if chosen else "",
            "confidence": ans["best_target"]["confidence"] if ans else None,
            "prob_none": probs.get("none", 0),
            "candidates": cands,
            "link_warranted": ans["link_warranted"]["noul"] if ans else None,
            "anchor_available": ans["anchor_available"]["noul"] if ans else None,
            "reader_stage": ans["reader_stage"]["choice"] if ans else "",
            "reader_stage_conf": ans["reader_stage"]["confidence"] if ans else None,
            "commercial": ans["commercial"]["noul"] if ans else None,
            "anchor": links[key]["anchor"] if key in links else "",
            "target_inbound_before": chosen["inbound_links"] if chosen else None,
            "score": link.get("final_score"),
            "model": resp.get("model", ""),
            "cost": (resp.get("usage") or {}).get("cost", 0) or 0,
            "tokens": (resp.get("usage") or {}).get("input_tokens", 0) or 0,
        }
        if compare_path:
            row["prev_choice"] = p["best_target"]["choice"] if p else ""
            row["prev_confidence"] = p["best_target"]["confidence"] if p else None
            row["same_as_prev"] = (None if not (p and ans) else
                                   p["best_target"]["choice"] == ans["best_target"]["choice"])
        rows.append(row)
    return rows


def write_csv(rows, compare):
    cols = ["id", "outcome_label", "source_url", "paragraph_no", "anchor_text",
            "jev_pick_url", "confidence", "prob_none"]
    for i in range(1, 5):
        cols += [f"cand{i}_url", f"cand{i}_prob", f"cand{i}_similarity"]
    cols += ["link_warranted", "anchor_available", "reader_stage",
             "commercial", "target_inbound_before", "final_score"]
    if compare:
        cols += ["prev_run_pick", "prev_run_confidence", "same_pick_as_prev_run"]
    cols += ["paragraph_text", "model"]

    f2 = lambda v: "" if v is None else f"{v:.2f}"
    with (OUT / "jev-decisions.csv").open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for r in rows:
            line = [r["id"], r["outcome_label"], r["source_url"],
                    r["paragraph_no"], r["anchor"], r["chosen_url"],
                    f2(r["confidence"]), f2(r["prob_none"])]
            for i in range(4):
                cd = r["candidates"][i] if i < len(r["candidates"]) else None
                line += ([cd["url"], f2(cd["prob"]), f2(cd["sim"])]
                         if cd else ["", "", ""])
            line += [f2(r["link_warranted"]), f2(r["anchor_available"]),
                     r["reader_stage"], f2(r["commercial"]),
                     "" if r["target_inbound_before"] is None else r["target_inbound_before"],
                     f2(r["score"])]
            if compare:
                line += [r["prev_choice"], f2(r["prev_confidence"]),
                         {True: "yes", False: "no"}.get(r["same_as_prev"], "")]
            line += [r["paragraph"], r["model"]]
            w.writerow(line)
    return cols


def write_xlsx(rows, compare):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    NAVY, PAPER = "15213B", "F2F4F7"
    fills = {
        "linked_apply": "DCE6FB", "linked_review": "F6E9CF",
        "no_anchor": "E9ECF1", "model_said_none": "F3DCE2",
    }
    head_font = Font(bold=True, color="FFFFFF", name="Calibri")
    head_fill = PatternFill("solid", fgColor=NAVY)

    wb = Workbook()

    def sheet(ws, headers, widths, data, wrap_cols=()):
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 22
        for rec, code in data:
            ws.append(rec)
            if code in fills:
                for cell in ws[ws.max_row]:
                    cell.fill = PatternFill("solid", fgColor=fills[code])
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = wdt
        for col in wrap_cols:
            for cell in ws[get_column_letter(col)][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    # Sheet 1: links to approve.
    ws = wb.active
    ws.title = "Links to approve"
    linked = [r for r in rows if r["outcome"].startswith("linked")]
    linked.sort(key=lambda r: -(r["score"] or 0))
    sheet(ws,
          ["Approve?", "Status", "Confidence", "Source page", "Para #",
           "Anchor text", "Links to", "Target inbound now", "Reader stage",
           "Paragraph"],
          [10, 14, 11, 42, 7, 28, 42, 10, 12, 90],
          [(["", "Apply" if r["outcome"] == "linked_apply" else "Review",
             round(r["confidence"], 2), path(r["source_url"]), r["paragraph_no"],
             r["anchor"], path(r["chosen_url"]), r["target_inbound_before"],
             r["reader_stage"], r["paragraph"]], r["outcome"]) for r in linked],
          wrap_cols=(10,))

    # Sheet 2: every decision.
    ws = wb.create_sheet("All decisions")
    hdr = ["#", "What happened", "Source page", "Para #", "Jev's pick",
           "Confidence", "P(none)", "Warranted", "Anchor avail.",
           "Reader stage", "Commercial", "Anchor text"]
    wid = [7, 34, 40, 7, 40, 11, 9, 10, 11, 12, 11, 26]
    if compare:
        hdr += ["Same pick as last run"]
        wid += [12]
    hdr += ["Paragraph"]
    wid += [90]
    data = []
    for r in rows:
        rec = [r["id"], r["outcome_label"], path(r["source_url"]),
               r["paragraph_no"],
               path(r["chosen_url"]) if r["chosen_url"] else "(none)",
               r["confidence"], r["prob_none"], r["link_warranted"],
               r["anchor_available"], r["reader_stage"], r["commercial"],
               r["anchor"]]
        if compare:
            rec.append({True: "yes", False: "no"}.get(r["same_as_prev"], ""))
        rec.append(r["paragraph"])
        data.append((rec, r["outcome"]))
    sheet(ws, hdr, wid, data, wrap_cols=(len(hdr),))
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=11):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.00"

    # Sheet 3: what happened, counted.
    ws = wb.create_sheet("Summary")
    counts = {}
    for r in rows:
        counts[r["outcome"]] = counts.get(r["outcome"], 0) + 1
    sheet(ws, ["What happened", "Decisions", "Share"], [48, 12, 10],
          [([OUTCOMES[k], v, v / len(rows)], k)
           for k, v in sorted(counts.items(), key=lambda kv: -kv[1])])
    for cell in ws["C"][1:]:
        cell.number_format = "0.0%"
    for s in wb.worksheets:
        s.sheet_properties.tabColor = NAVY
    wb.save(OUT / "jev-decisions.xlsx")


def summarise(rows, compare):
    paired = [r for r in rows if compare and r.get("same_as_prev") is not None]
    s = {
        "site": urlparse(rows[0]["source_url"]).netloc if rows else "",
        "date": date.today().isoformat(),
        "model": next((r["model"] for r in rows if r["model"]), ""),
        "pages": len(json.loads((DATA / "pages.json").read_text())),
        "decisions": len(rows),
        "cost": round(sum(r["cost"] for r in rows), 4),
        "tokens": sum(r["tokens"] for r in rows),
        "thresholds": CFG["thresholds"],
        "caps": CFG["caps"],
        "outcomes": OUTCOMES,
    }
    if paired:
        hi = [r for r in paired if (r["prev_confidence"] or 0) >= CFG["thresholds"]["auto_approve_confidence"]]
        s["consistency"] = {
            "paired": len(paired),
            "same_pick": sum(r["same_as_prev"] for r in paired),
            "high_conf": len(hi),
            "high_conf_same": sum(r["same_as_prev"] for r in hi),
            "median_conf_shift": round(statistics.median(
                abs(r["confidence"] - r["prev_confidence"]) for r in paired), 3),
        }
    return s


def write_dashboard(rows, summary):
    # Compact payload: short keys, target pages de-duplicated into one table.
    urls, uidx = [], {}

    def u(url, title):
        if url not in uidx:
            uidx[url] = len(urls)
            urls.append([path(url), title])
        return uidx[url]

    rnd = lambda v: None if v is None else round(v, 3)
    D = []
    for r in rows:
        D.append({
            "i": r["id"], "o": r["outcome"],
            "s": u(r["source_url"], r["source_title"]), "p": r["paragraph_no"],
            "t": r["paragraph"], "a": r["anchor"],
            "c": rnd(r["confidence"]), "n": rnd(r["prob_none"]),
            "k": [[u(cd["url"], cd["title"]), rnd(cd["prob"]), rnd(cd["sim"]), cd["inbound"]]
                  for cd in r["candidates"]],
            "j": r["jev_choice"],
            "w": rnd(r["link_warranted"]), "v": rnd(r["anchor_available"]),
            "r": r["reader_stage"], "m": rnd(r["commercial"]),
            "x": r.get("same_as_prev"), "y": r.get("prev_choice"),
        })
    payload = json.dumps({"summary": summary, "urls": urls, "rows": D},
                         separators=(",", ":")).replace("</", "<\\/")
    html = TEMPLATE.read_text().replace("/*__DATA__*/null", payload)
    (OUT / "jev-dashboard.html").write_text(html)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--compare", help="judgements.json from an earlier run")
    args = ap.parse_args()

    OUT.mkdir(exist_ok=True)
    rows = build_rows(args.compare)
    compare = bool(args.compare)
    write_csv(rows, compare)
    write_xlsx(rows, compare)
    summary = summarise(rows, compare)
    write_dashboard(rows, summary)

    counts = {}
    for r in rows:
        counts[r["outcome"]] = counts.get(r["outcome"], 0) + 1
    print(f"{len(rows):,} decisions, model {summary['model']}, ${summary['cost']:.4f}")
    for k, v in sorted(counts.items(), key=lambda kv: -kv[1]):
        print(f"   {OUTCOMES[k]:45} {v:,}")
    if "consistency" in summary:
        c = summary["consistency"]
        print(f"same pick as earlier run: {c['same_pick']}/{c['paired']}  "
              f"(high-confidence: {c['high_conf_same']}/{c['high_conf']})")
    for f in ("jev-decisions.csv", "jev-decisions.xlsx", "jev-dashboard.html"):
        print(f"wrote {OUT / f}")


if __name__ == "__main__":
    main()
